import io
import uuid
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings as app_settings
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.knowledge import Document, DocumentChunk, DocumentStatus
from app.rag import embedder, vectordb
from app.rag.chunker import CHUNK_SIZE_CHARS, OVERLAP_CHARS, chunk_pages
from app.rag.parser import Page, ParsingError, parse_file
from app.services.storage import get_uploads_dir

# --- fixtures ---------------------------------------------------------------


@pytest.fixture()
def db_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    testing_session = sessionmaker(
        bind=engine,
        autoflush=False,
        autocommit=False,
        expire_on_commit=False,
    )
    session = testing_session()
    yield session
    session.close()


@pytest.fixture()
def uploads_dir(tmp_path: Path) -> Path:
    path = tmp_path / "uploads"
    path.mkdir()
    return path


@pytest.fixture()
def chroma_dir(tmp_path: Path, monkeypatch) -> Path:
    path = tmp_path / "chroma_db"
    monkeypatch.setattr(app_settings, "chroma_path", str(path))
    return path


@pytest.fixture()
def fake_embedder(monkeypatch) -> None:
    def _encode(texts):
        return [
            [float(len(text)), float(sum(ord(c) for c in text) % 101), 1.0]
            for text in texts
        ]

    monkeypatch.setattr(embedder, "encode_texts", _encode)


@pytest.fixture()
def client(db_session, uploads_dir: Path, chroma_dir: Path, fake_embedder, rbac_seed):
    def _override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = _override_get_db
    app.dependency_overrides[get_uploads_dir] = lambda: uploads_dir

    with TestClient(app, base_url="https://testserver") as test_client:
        response = test_client.post(
            "/api/auth/register",
            json={"email": "admin@orbit.test", "password": "password123"},
        )
        assert response.status_code == 201, response.text
        yield test_client

    app.dependency_overrides.clear()


# --- helpers ----------------------------------------------------------------


def upload_file(
    client: TestClient,
    filename: str,
    content: bytes,
    content_type: str = "text/plain",
) -> dict[str, Any]:
    response = client.post(
        "/api/documents/upload",
        files={"file": (filename, content, content_type)},
    )
    assert response.status_code == 201, response.text
    return response.json()


def index_document(client: TestClient, document_id: str) -> dict[str, Any]:
    response = client.post(f"/api/documents/{document_id}/index")
    assert response.status_code == 200, response.text
    return response.json()


def make_pdf_bytes(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    payload = doc.tobytes()
    doc.close()
    return payload


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- parser -----------------------------------------------------------------


def test_parse_pdf_extracts_pages_with_numbers(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    path.write_bytes(make_pdf_bytes(["First page content", "Second page content"]))

    pages = parse_file(path)

    assert [p.page_number for p in pages] == [1, 2]
    assert "First page content" in pages[0].text
    assert "Second page content" in pages[1].text


def test_parse_txt_yields_single_page(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("line one\nline two", encoding="utf-8")

    pages = parse_file(path)

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].text == "line one\nline two"


def test_parse_md_yields_single_page(tmp_path: Path) -> None:
    path = tmp_path / "readme.md"
    path.write_text("# Title\n\nBody", encoding="utf-8")

    pages = parse_file(path)

    assert len(pages) == 1
    assert "# Title" in pages[0].text


def test_parse_docx_yields_single_page(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    path.write_bytes(make_docx_bytes(["Alpha paragraph", "Beta paragraph"]))

    pages = parse_file(path)

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "Alpha paragraph" in pages[0].text
    assert "Beta paragraph" in pages[0].text


def test_parse_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(ParsingError):
        parse_file(tmp_path / "nope.txt")


def test_parse_corrupt_pdf_raises(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 this is not really a pdf")

    with pytest.raises(ParsingError):
        parse_file(path)


def test_parse_unsupported_suffix_raises(tmp_path: Path) -> None:
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04")

    with pytest.raises(ParsingError):
        parse_file(path)


# --- chunker ----------------------------------------------------------------


def test_chunker_constants_match_spec() -> None:
    assert CHUNK_SIZE_CHARS == 2000  # 500 tokens * 4 chars
    assert OVERLAP_CHARS == 400  # 100 tokens * 4 chars


def test_chunk_short_page_returns_single_ordered_chunk() -> None:
    chunks = chunk_pages([Page(page_number=1, text="A short paragraph.")])

    assert len(chunks) == 1
    assert chunks[0].chunk_index == 0
    assert chunks[0].page == 1
    assert chunks[0].text.strip() == "A short paragraph."


def test_chunk_long_text_splits_within_size_with_overlap() -> None:
    text = " ".join(f"Sentence number {i} has unique words." for i in range(400))

    chunks = chunk_pages([Page(page_number=1, text=text)])

    assert len(chunks) > 1
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert all(len(c.text) <= CHUNK_SIZE_CHARS for c in chunks)
    # Coverage: first chunk opens the text, last chunk closes it.
    assert chunks[0].text.startswith(text[:20])
    assert text[-20:] in chunks[-1].text
    # Overlap: each chunk after the first opens with a tail of its predecessor.
    assert chunks[0].text[-250:] in chunks[1].text


def test_chunk_preserves_page_metadata_and_global_order() -> None:
    long_page_one = " ".join(f"Page one sentence {i}." for i in range(300))
    pages = [
        Page(page_number=1, text=long_page_one),
        Page(page_number=2, text="Page two text."),
        Page(page_number=3, text="Page three text."),
    ]

    chunks = chunk_pages(pages)

    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert chunks[0].page == 1
    assert chunks[-1].page == 3
    page_two = [c for c in chunks if c.page == 2]
    assert len(page_two) == 1
    assert page_two[0].text == "Page two text."
    # Page numbers are non-decreasing in chunk order.
    assert [c.page for c in chunks] == sorted(c.page for c in chunks)


def test_chunk_skips_empty_pages() -> None:
    chunks = chunk_pages(
        [Page(page_number=1, text="   \n "), Page(page_number=2, text="hi")]
    )

    assert len(chunks) == 1
    assert chunks[0].page == 2


# --- vector store -----------------------------------------------------------


def test_vectordb_add_and_delete_by_document(tmp_path: Path) -> None:
    collection = vectordb.get_collection(path=str(tmp_path / "chroma"))

    vectordb.upsert_chunks(
        collection,
        ids=["doc-a:0", "doc-a:1", "doc-b:0"],
        embeddings=[[1.0, 0.0, 0.0], [0.0, 1.0, 0.0], [0.5, 0.5, 0.0]],
        documents=["chunk a0", "chunk a1", "chunk b0"],
        metadatas=[
            {"document_id": "doc-a", "page": 1, "chunk_index": 0, "filename": "a.txt"},
            {"document_id": "doc-a", "page": 2, "chunk_index": 1, "filename": "a.txt"},
            {"document_id": "doc-b", "page": 1, "chunk_index": 0, "filename": "b.txt"},
        ],
    )
    assert vectordb.count_document_vectors(collection, "doc-a") == 2

    vectordb.delete_document_vectors(collection, "doc-a")

    assert vectordb.count_document_vectors(collection, "doc-a") == 0
    assert vectordb.count_document_vectors(collection, "doc-b") == 1


def test_vectordb_persists_to_disk(tmp_path: Path) -> None:
    import chromadb

    path = tmp_path / "chroma_db"
    collection = vectordb.get_collection(path=str(path))
    vectordb.upsert_chunks(
        collection,
        ids=["d:0"],
        embeddings=[[0.1, 0.2, 0.3]],
        documents=["persisted"],
        metadatas=[
            {"document_id": "d", "page": 1, "chunk_index": 0, "filename": "f.txt"}
        ],
    )

    fresh = chromadb.PersistentClient(path=str(path)).get_or_create_collection(
        name=vectordb.COLLECTION_NAME
    )
    assert fresh.count() == 1
    assert path.exists()


# --- embedder ---------------------------------------------------------------


def test_encode_empty_list_returns_empty() -> None:
    assert embedder.encode_texts([]) == []


def test_embedder_loads_real_model_once() -> None:
    try:
        model = embedder.get_model()
        vectors = embedder.encode_texts(["orbit knowledge base"])
    except Exception as exc:  # pragma: no cover - environment dependent
        pytest.skip(f"embedding model unavailable: {exc}")

    assert model is embedder.get_model()
    assert len(vectors) == 1
    assert len(vectors[0]) == 384
    magnitude = sum(value * value for value in vectors[0]) ** 0.5
    assert magnitude == pytest.approx(1.0, abs=1e-3)


# --- indexer (service level) ------------------------------------------------


def _seed_document(
    db_session, uploads_dir: Path, name: str, content: bytes
) -> Document:
    stored = f"{uuid.uuid4()}{Path(name).suffix}"
    (uploads_dir / stored).write_bytes(content)
    document = Document(
        original_filename=name,
        stored_filename=stored,
        file_extension=Path(name).suffix.lstrip("."),
        mime_type="text/plain",
        file_size=len(content),
    )
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)
    return document


def test_index_document_service_writes_chunks_and_vectors(
    db_session, uploads_dir: Path, chroma_dir: Path, fake_embedder
) -> None:
    from app.rag.indexer import index_document as run_indexing

    text = " ".join(f"Indexed sentence {i} with content." for i in range(250))
    document = _seed_document(db_session, uploads_dir, "long.txt", text.encode())

    result = run_indexing(db_session, document, uploads_dir)

    assert result.status == DocumentStatus.INDEXED
    assert result.progress == 100
    assert result.chunk_count >= 2

    db_session.refresh(document)
    assert document.status == DocumentStatus.INDEXED
    assert document.chunk_count == result.chunk_count
    assert document.index_error is None

    rows = db_session.scalars(
        select(DocumentChunk).where(DocumentChunk.document_id == document.id)
    ).all()
    assert len(rows) == result.chunk_count
    assert [r.chunk_index for r in rows] == sorted(r.chunk_index for r in rows)
    assert all(r.page == 1 for r in rows)

    collection = vectordb.get_collection()
    assert (
        vectordb.count_document_vectors(collection, document.id) == result.chunk_count
    )


def test_index_document_service_handles_missing_file(
    db_session, uploads_dir: Path, chroma_dir: Path, fake_embedder
) -> None:
    from app.rag.indexer import index_document as run_indexing

    document = _seed_document(db_session, uploads_dir, "gone.txt", b"hello")
    (uploads_dir / document.stored_filename).unlink()

    result = run_indexing(db_session, document, uploads_dir)

    assert result.status == DocumentStatus.FAILED
    assert result.chunk_count == 0
    assert result.error
    db_session.refresh(document)
    assert document.status == DocumentStatus.FAILED
    assert document.index_error


def test_index_document_service_handles_embedder_failure(
    db_session, uploads_dir: Path, chroma_dir: Path, monkeypatch
) -> None:
    from app.rag.indexer import index_document as run_indexing

    def _explode(_texts):
        raise RuntimeError("embedding backend down")

    monkeypatch.setattr(embedder, "encode_texts", _explode)
    document = _seed_document(db_session, uploads_dir, "text.txt", b"some content")

    result = run_indexing(db_session, document, uploads_dir)

    assert result.status == DocumentStatus.FAILED
    assert "embedding backend down" in (result.error or "")
    db_session.refresh(document)
    assert document.status == DocumentStatus.FAILED
    assert document.chunk_count == 0


# --- API --------------------------------------------------------------------


def test_index_endpoint_indexes_and_reports_progress(client: TestClient) -> None:
    content = " ".join(f"Document sentence {i} goes here." for i in range(250))
    document = upload_file(client, "report.txt", content.encode())

    payload = index_document(client, document["id"])

    assert payload["document_id"] == document["id"]
    assert payload["status"] == "indexed"
    assert payload["progress"] == 100
    assert payload["chunk_count"] >= 2
    assert payload["error"] is None


def test_list_documents_includes_status_and_chunk_count(client: TestClient) -> None:
    content = " ".join(f"Word {i} of the corpus." for i in range(250))
    document = upload_file(client, "corpus.txt", content.encode())

    before = client.get("/api/documents").json()["items"][0]
    assert before["status"] == "uploaded"
    assert before["chunk_count"] == 0
    assert before["index_error"] is None

    index_document(client, document["id"])

    after = client.get("/api/documents").json()["items"][0]
    assert after["status"] == "indexed"
    assert after["chunk_count"] >= 2
    assert after["index_error"] is None


def test_index_endpoint_indexes_real_pdf_with_page_metadata(
    client: TestClient, db_session
) -> None:
    pdf = make_pdf_bytes(
        ["Alpha page discusses contracts.", "Beta page discusses risks."]
    )
    document = upload_file(client, "contracts.pdf", pdf, "application/pdf")

    payload = index_document(client, document["id"])

    assert payload["status"] == "indexed"
    rows = db_session.scalars(
        select(DocumentChunk).where(DocumentChunk.document_id == document["id"])
    ).all()
    assert {row.page for row in rows} == {1, 2}

    collection = vectordb.get_collection()
    stored = collection.get(
        where={"document_id": document["id"]}, include=["metadatas"]
    )
    assert len(stored["ids"]) == payload["chunk_count"]
    for metadata in stored["metadatas"]:
        assert metadata["document_id"] == document["id"]
        assert metadata["filename"] == "contracts.pdf"
        assert metadata["page"] in (1, 2)
        assert isinstance(metadata["chunk_index"], int)


def test_index_endpoint_indexes_docx(client: TestClient) -> None:
    docx = make_docx_bytes(["Onboarding guide paragraph one.", "Paragraph two."])
    document = upload_file(
        client,
        "guide.docx",
        docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    payload = index_document(client, document["id"])

    assert payload["status"] == "indexed"
    assert payload["chunk_count"] >= 1


def test_index_endpoint_reindex_is_idempotent(client: TestClient, db_session) -> None:
    content = " ".join(f"Repeatable sentence {i}." for i in range(200))
    document = upload_file(client, "repeat.txt", content.encode())

    first = index_document(client, document["id"])
    second = index_document(client, document["id"])

    assert first["status"] == second["status"] == "indexed"
    assert first["chunk_count"] == second["chunk_count"]

    rows = db_session.scalars(
        select(DocumentChunk).where(DocumentChunk.document_id == document["id"])
    ).all()
    assert len(rows) == first["chunk_count"]

    collection = vectordb.get_collection()
    assert (
        vectordb.count_document_vectors(collection, document["id"])
        == first["chunk_count"]
    )


def test_index_endpoint_corrupt_pdf_fails_gracefully(client: TestClient) -> None:
    document = upload_file(
        client, "broken.pdf", b"%PDF-1.4 not really", "application/pdf"
    )

    payload = index_document(client, document["id"])

    assert payload["status"] == "failed"
    assert payload["chunk_count"] == 0
    assert payload["error"]
    assert payload["progress"] < 100

    listed = client.get("/api/documents").json()["items"][0]
    assert listed["status"] == "failed"
    assert listed["chunk_count"] == 0
    assert listed["index_error"]


def test_index_endpoint_unknown_document_returns_404(client: TestClient) -> None:
    response = client.post("/api/documents/missing-id/index")
    assert response.status_code == 404


def test_index_endpoint_missing_stored_file_fails_gracefully(
    client: TestClient, uploads_dir: Path
) -> None:
    document = upload_file(client, "vanished.txt", b"content")
    for path in uploads_dir.iterdir():
        path.unlink()

    payload = index_document(client, document["id"])

    assert payload["status"] == "failed"
    assert payload["error"]


def test_index_endpoint_after_delete_returns_404(client: TestClient) -> None:
    document = upload_file(client, "temp.txt", b"content")
    assert client.delete(f"/api/documents/{document['id']}").status_code == 200

    response = client.post(f"/api/documents/{document['id']}/index")
    assert response.status_code == 404
